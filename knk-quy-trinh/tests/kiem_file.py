#!/usr/bin/env python3
"""Doc file .docx va .xlsx do ung dung xuat, in ra JSON de kiem thu Giai doan 5 so sanh.

    python3 tests/kiem_file.py <file.docx|-> <file.xlsx|-> <thu_muc_anh> [tien_to_anh]

.docx: doan van, bang, cac run to vang, thong so trang, dau trang; mo bang LibreOffice,
       xuat PDF, dem so trang va chup anh tung trang vao thu muc anh.
.xlsx: cong thuc va gia tri tinh san cua moi o; xoa gia tri tinh san cua cac o cong thuc
       roi cho LibreOffice tinh lai, doc gia tri moi tinh. Hai bo gia tri phai khop nhau.

Can: python-docx, openpyxl, pymupdf, LibreOffice Writer va Calc (soffice).
"""
import json, os, re, shutil, subprocess, sys, tempfile, zipfile

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def soffice(args, cwd):
    exe = shutil.which('soffice') or shutil.which('libreoffice')
    if not exe:
        return False
    prof = 'file://' + os.path.join(cwd, 'lo_profile')
    r = subprocess.run([exe, '--headless', '--norestore', '-env:UserInstallation=' + prof] + args,
                       cwd=cwd, capture_output=True, text=True, timeout=240)
    return r.returncode == 0


def doc_docx(path, anh, tien_to):
    import docx
    d = docx.Document(path)
    out = {'paras': [], 'tables': [], 'highlights': [], 'sect': {}, 'parts': []}
    body = d.element.body
    for el in body.iterchildren():
        if el.tag == W + 'p':
            out['paras'].append(''.join(t.text or '' for t in el.iter(W + 't')))
        elif el.tag == W + 'tbl':
            rows = []
            for tr in el.iter(W + 'tr'):
                rows.append([''.join(t.text or '' for t in tc.iter(W + 't')) for tc in tr.iter(W + 'tc')])
            out['tables'].append(rows)
    for r in body.iter(W + 'r'):
        if r.find(W + 'rPr/' + W + 'highlight') is not None:
            out['highlights'].append(''.join(t.text or '' for t in r.iter(W + 't')))
    s = d.sections[0]
    out['sect'] = {'w': s.page_width.twips if s.page_width else None, 'h': s.page_height.twips if s.page_height else None,
                   'top': s.top_margin.twips, 'bottom': s.bottom_margin.twips, 'left': s.left_margin.twips,
                   'right': s.right_margin.twips, 'titlePg': s.different_first_page_header_footer}
    out['header'] = ''.join(''.join(x.itertext()) for x in s.header._element.iter(W + 'instrText'))
    z = zipfile.ZipFile(path)
    out['parts'] = sorted(z.namelist())
    xml = z.read('word/document.xml').decode('utf-8')
    out['fonts'] = sorted(set(re.findall(r'w:ascii="([^"]+)"', z.read('word/styles.xml').decode('utf-8'))))
    out['subscripts'] = xml.count('<w:vertAlign w:val="subscript"/>')
    # mo bang LibreOffice Writer
    tmp = tempfile.mkdtemp(prefix='kf_')
    shutil.copy(path, os.path.join(tmp, 'm.docx'))
    out['pdf'] = None
    if soffice(['--convert-to', 'pdf', '--outdir', tmp, 'm.docx'], tmp) and os.path.exists(os.path.join(tmp, 'm.pdf')):
        import pymupdf
        pdf = pymupdf.open(os.path.join(tmp, 'm.pdf'))
        out['pdf'] = {'pages': pdf.page_count}
        if anh:
            os.makedirs(anh, exist_ok=True)
            for i, pg in enumerate(pdf):
                pg.get_pixmap(dpi=80).save(os.path.join(anh, '%s_trang_%d.png' % (tien_to, i + 1)))
    shutil.rmtree(tmp, ignore_errors=True)
    return out


def doc_xlsx(path):
    import openpyxl
    out = {'sheets': [], 'cells': {}, 'calc': {}}
    wb = openpyxl.load_workbook(path)
    out['sheets'] = wb.sheetnames
    wv = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = wv[ws.title][c.coordinate].value
                if isinstance(c.value, str) and c.value.startswith('='):
                    cells[c.coordinate] = {'f': c.value, 'v': v}
                else:
                    cells[c.coordinate] = {'v': c.value}
        out['cells'][ws.title] = cells
    # bo gia tri tinh san cua o cong thuc, de LibreOffice phai tu tinh
    tmp = tempfile.mkdtemp(prefix='kf_')
    src = zipfile.ZipFile(path)
    dst = zipfile.ZipFile(os.path.join(tmp, 'tho.xlsx'), 'w')
    for it in src.infolist():
        data = src.read(it.filename)
        if it.filename.startswith('xl/worksheets/'):
            data = re.sub(rb'(<f>[^<]*</f>)<v>[^<]*</v>', rb'\1', data)
        dst.writestr(it.filename, data)
    dst.close()
    os.makedirs(os.path.join(tmp, 'ra'))
    out['calc'] = None
    if soffice(['--convert-to', 'xlsx', '--outdir', os.path.join(tmp, 'ra'), 'tho.xlsx'], tmp):
        wc = openpyxl.load_workbook(os.path.join(tmp, 'ra', 'tho.xlsx'), data_only=True)
        out['calc'] = {}
        for ws in wb.worksheets:
            out['calc'][ws.title] = {ref: wc[ws.title][ref].value for ref, c in out['cells'][ws.title].items() if 'f' in c}
    shutil.rmtree(tmp, ignore_errors=True)
    return out


def main():
    docx_p, xlsx_p, anh = sys.argv[1], sys.argv[2], sys.argv[3]
    tien_to = sys.argv[4] if len(sys.argv) > 4 else 'mau06'
    kq = {}
    if docx_p != '-':
        kq['docx'] = doc_docx(docx_p, anh, tien_to)
    if xlsx_p != '-':
        kq['xlsx'] = doc_xlsx(xlsx_p)
    json.dump(kq, sys.stdout, ensure_ascii=False, default=str)


if __name__ == '__main__':
    main()
